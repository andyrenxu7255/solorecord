from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from .config import get_settings
from .db import get_db
from .repository import action_items_with_evidence, build_quality_report
from .utils import new_id, now_iso


def create_export(meeting_id: str, export_format: str) -> dict:
    export_format = export_format.lower()
    if export_format not in {"markdown", "json", "srt", "docx", "pdf"}:
        raise ValueError("Unsupported export format")
    settings = get_settings()
    export_dir = settings.storage_dir / "exports" / meeting_id
    export_dir.mkdir(parents=True, exist_ok=True)
    file_name = f"{meeting_id}.{_extension(export_format)}"
    path = export_dir / file_name
    meeting, segments, actions, audio_rows = _load_meeting(meeting_id)
    quality_report = build_quality_report(
        segments,
        actions,
        audio_rows,
        meeting.get("summary", ""),
        meeting.get("role_notes", ""),
    )
    enriched_actions = action_items_with_evidence(actions, quality_report)
    if export_format == "markdown":
        path.write_text(_markdown(meeting, segments, enriched_actions), encoding="utf-8")
    elif export_format == "json":
        import json

        path.write_text(
            json.dumps(
                {
                    "meeting": meeting,
                    "segments": segments,
                    "actions": enriched_actions,
                    "qualityReport": quality_report,
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
    elif export_format == "srt":
        path.write_text(_srt(segments), encoding="utf-8")
    elif export_format == "docx":
        _docx(path, meeting, segments, enriched_actions)
    elif export_format == "pdf":
        _pdf(path, meeting, segments, enriched_actions)
    export_id = new_id("exp")
    with get_db() as db:
        db.execute(
            """
            INSERT INTO exports (id, meeting_id, format, storage_path, created_at)
            VALUES (?, ?, ?, ?, ?)
            """,
            (export_id, meeting_id, export_format, str(path), now_iso()),
        )
    return {"id": export_id, "format": export_format, "path": str(path), "file_name": file_name}


def _load_meeting(meeting_id: str) -> tuple[dict, list[dict], list[dict], list[dict]]:
    with get_db() as db:
        meeting = db.execute("SELECT * FROM meetings WHERE id = ?", (meeting_id,)).fetchone()
        if not meeting:
            raise ValueError("Meeting not found")
        segments = db.execute(
            "SELECT * FROM transcript_segments WHERE meeting_id = ? ORDER BY start_ms",
            (meeting_id,),
        ).fetchall()
        actions = db.execute(
            "SELECT * FROM action_items WHERE meeting_id = ? ORDER BY created_at",
            (meeting_id,),
        ).fetchall()
        audio_rows = db.execute(
            "SELECT * FROM audio_segments WHERE meeting_id = ? ORDER BY segment_no",
            (meeting_id,),
        ).fetchall()
    return (
        dict(meeting),
        [dict(row) for row in segments],
        [dict(row) for row in actions],
        [dict(row) for row in audio_rows],
    )


def _extension(export_format: str) -> str:
    return {"markdown": "md", "json": "json", "srt": "srt", "docx": "docx", "pdf": "pdf"}[export_format]


def _markdown(meeting: dict, segments: list[dict], actions: list[dict]) -> str:
    meeting_actions, review_actions = _split_review_actions(actions)
    lines = [
        f"# {meeting['title']}",
        "",
        "## 会议纪要",
        "",
        meeting.get("summary") or "暂无纪要",
        "",
        "## 分角色整理",
        "",
        meeting.get("role_notes") or "暂无分角色整理",
        "",
        "## 转写",
        "",
    ]
    for segment in segments:
        lines.append(f"- [{_time(segment['start_ms'])}] **{segment['display_name']}**：{segment['text']}")
    lines.extend(["", "## 待办", ""])
    for item in meeting_actions:
        lines.append(f"- [{item['status']}] {item['owner']}：{item['task']} {item['due']}")
    if not meeting_actions:
        lines.append("- 暂无可督办待办")
    if review_actions:
        lines.extend(["", "## 系统复核提醒", ""])
        for item in review_actions:
            reason = item.get("evidenceReason") or "不是可自动督办的会议待办"
            lines.append(f"- {item['owner']}：{item['task']}（{reason}）")
    return "\n".join(lines) + "\n"


def _srt(segments: list[dict]) -> str:
    blocks = []
    for index, segment in enumerate(segments, start=1):
        blocks.append(
            f"{index}\n{_srt_time(segment['start_ms'])} --> {_srt_time(segment['end_ms'])}\n"
            f"{segment['display_name']}：{segment['text']}\n"
        )
    return "\n".join(blocks)


def _docx(path: Path, meeting: dict, segments: list[dict], actions: list[dict]) -> None:
    meeting_actions, review_actions = _split_review_actions(actions)
    doc = Document()
    doc.add_heading(meeting["title"], level=1)
    doc.add_heading("会议纪要", level=2)
    doc.add_paragraph(meeting.get("summary") or "暂无纪要")
    doc.add_heading("分角色整理", level=2)
    doc.add_paragraph(meeting.get("role_notes") or "暂无分角色整理")
    doc.add_heading("转写", level=2)
    for segment in segments:
        doc.add_paragraph(f"[{_time(segment['start_ms'])}] {segment['display_name']}：{segment['text']}")
    doc.add_heading("待办", level=2)
    if not meeting_actions:
        doc.add_paragraph("暂无可督办待办")
    for item in meeting_actions:
        doc.add_paragraph(f"{item['owner']}：{item['task']} {item['due']} [{item['status']}]")
    if review_actions:
        doc.add_heading("系统复核提醒", level=2)
        for item in review_actions:
            reason = item.get("evidenceReason") or "不是可自动督办的会议待办"
            doc.add_paragraph(f"{item['owner']}：{item['task']}（{reason}）")
    doc.save(path)


def _pdf(path: Path, meeting: dict, segments: list[dict], actions: list[dict]) -> None:
    meeting_actions, review_actions = _split_review_actions(actions)
    pdf = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    y = height - 40
    for line in [
        meeting["title"],
        "会议纪要",
        meeting.get("summary") or "暂无纪要",
        "分角色整理",
        meeting.get("role_notes") or "暂无分角色整理",
        "转写",
    ]:
        pdf.drawString(40, y, line[:90])
        y -= 22
    for segment in segments:
        text = f"[{_time(segment['start_ms'])}] {segment['display_name']}: {segment['text']}"
        pdf.drawString(40, y, text[:110])
        y -= 18
        if y < 60:
            pdf.showPage()
            y = height - 40
    pdf.drawString(40, y, "待办")
    y -= 22
    if not meeting_actions:
        pdf.drawString(40, y, "暂无可督办待办")
        y -= 18
    for item in meeting_actions:
        pdf.drawString(40, y, f"{item['owner']}: {item['task']} {item['due']} [{item['status']}]"[:110])
        y -= 18
        if y < 60:
            pdf.showPage()
            y = height - 40
    if review_actions:
        pdf.drawString(40, y, "系统复核提醒")
        y -= 22
        for item in review_actions:
            reason = item.get("evidenceReason") or "不是可自动督办的会议待办"
            pdf.drawString(40, y, f"{item['owner']}: {item['task']} ({reason})"[:110])
            y -= 18
            if y < 60:
                pdf.showPage()
                y = height - 40
    pdf.save()


def _split_review_actions(actions: list[dict]) -> tuple[list[dict], list[dict]]:
    meeting_actions = []
    review_actions = []
    for item in actions:
        if _is_review_action(item):
            review_actions.append(item)
        else:
            meeting_actions.append(item)
    return meeting_actions, review_actions


def _is_review_action(item: dict) -> bool:
    return bool(
        item.get("reviewOnly")
        or item.get("review_only")
        or item.get("actionKind") == "system_review"
        or item.get("action_kind") == "system_review"
        or item.get("evidenceStatus") == "system_review"
    )


def _time(ms: int) -> str:
    seconds = max(0, int(ms / 1000))
    return f"{seconds // 3600:02d}:{seconds % 3600 // 60:02d}:{seconds % 60:02d}"


def _srt_time(ms: int) -> str:
    hours = ms // 3_600_000
    minutes = (ms % 3_600_000) // 60_000
    seconds = (ms % 60_000) // 1000
    millis = ms % 1000
    return f"{hours:02d}:{minutes:02d}:{seconds:02d},{millis:03d}"
